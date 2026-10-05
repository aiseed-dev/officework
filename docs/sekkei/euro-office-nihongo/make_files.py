"""Write the small documents that show how Euro-Office (and ONLYOFFICE,
whose x2t was used to check them) handles Japanese.

    .venv/bin/python docs/sekkei/euro-office-nihongo/make_files.py docs/sekkei/euro-office-nihongo

The docx files use IPAex Mincho at 10.5pt on a line exactly 40
full-width characters wide (420pt), so where a line breaks can be read
from a PDF. Print them with tools/oo_pdf.py (ONLYOFFICE) and
tools/lo_pdf.py (LibreOffice), and compare with Word and Excel on the Mac.
"""
import datetime
import sys
import zipfile

import openpyxl
from openpyxl.styles import Alignment, Font

OUT = sys.argv[1]
W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
FONT = '<w:rFonts w:ascii="IPAexMincho" w:hAnsi="IPAexMincho" w:eastAsia="IPAexMincho"/><w:sz w:val="21"/><w:szCs w:val="21"/>'


def run(text, extra=""):
    return f'<w:r><w:rPr>{FONT}{extra}</w:rPr><w:t xml:space="preserve">{text}</w:t></w:r>'


def para(text, ppr=""):
    return f'<w:p><w:pPr><w:spacing w:after="120"/>{ppr}</w:pPr>{run(text)}</w:p>'


def sect(extra=""):
    return ('<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" w:right="1753" w:bottom="1440" '
            f'w:left="1753" w:header="720" w:footer="720" w:gutter="0"/>{extra}</w:sectPr>')


def docx(name, body):
    ct = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
          '<Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" '
          'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
    doc = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {W}><w:body>{body}</w:body></w:document>'
    with zipfile.ZipFile(f"{OUT}/{name}", "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", doc)


# Line breaking: the 41st character falls at the start of line 2
fill = "漢" * 39 + "字"
body = "".join(para(fill + c + "以下続く文章です。") for c in "。、」』）】！？・：ー々ゝっゃッャ")
body += para("漢" * 39 + "「括弧の中」です。")
docx("01-kinsoku.docx", body + sect())
# Kana runs: Word fills the first line
docx("02-kana-line-break.docx",
     para("漢字" * 18 + "ですあいうえおかきくけこさしすせそたちつてと")
     + para("東京" * 19 + "にすんでいます。わたしはがくせいです。") + sect())
# Ruby
ruby = ('<w:p>' + run("氏名：") + f'<w:r><w:rPr>{FONT}</w:rPr><w:ruby><w:rubyPr><w:rubyAlign w:val="distributeSpace"/><w:hps w:val="10"/>'
        '<w:hpsRaise w:val="20"/><w:hpsBaseText w:val="21"/><w:lid w:val="ja-JP"/></w:rubyPr>'
        '<w:rt><w:r><w:rPr><w:rFonts w:ascii="IPAexMincho" w:hAnsi="IPAexMincho" w:eastAsia="IPAexMincho"/><w:sz w:val="10"/></w:rPr><w:t>やまだ</w:t></w:r></w:rt>'
        f'<w:rubyBase><w:r><w:rPr>{FONT}</w:rPr><w:t>山田</w:t></w:r></w:rubyBase></w:ruby></w:r>' + run("太郎") + '</w:p>')
docx("03-ruby.docx", ruby + sect())
# Vertical writing: a section, and a table cell
docx("04-vertical-section.docx",
     para("縦書きの文章です。数字は２０２６年、英字はＡＢＣとabcです。「括弧」も縦になります。")
     + para("二行目の段落です。") + sect('<w:textDirection w:val="tbRl"/>'))
cell = ('<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/><w:tblBorders><w:top w:val="single" w:sz="4"/><w:left w:val="single" w:sz="4"/>'
        '<w:bottom w:val="single" w:sz="4"/><w:right w:val="single" w:sz="4"/></w:tblBorders></w:tblPr><w:tblGrid><w:gridCol w:w="800"/></w:tblGrid>'
        '<w:tr><w:trPr><w:trHeight w:val="3000"/></w:trPr><w:tc><w:tcPr><w:tcW w:w="800" w:type="dxa"/><w:textDirection w:val="tbRl"/></w:tcPr>'
        '<w:p>' + run("縦書きＡＢ「題」") + '</w:p></w:tc></w:tr></w:tbl><w:p/>')
docx("05-vertical-cell.docx", cell + sect())
# Japanese paragraph and character properties that a round trip should keep
body = (
    '<w:p>' + run("傍点", '<w:em w:val="dot"/>') + run("と") + run("傍点", '<w:em w:val="comma"/>') + '</w:p>'
    '<w:p>' + run("令和") + run("12", '<w:eastAsianLayout w:id="1" w:vert="1"/>') + run("年") + '</w:p>'
    '<w:p>' + run("本文") + run("割注の例", '<w:eastAsianLayout w:id="2" w:combine="1" w:combineBrackets="round"/>') + '</w:p>'
    '<w:p>' + run("氏名", '<w:fitText w:val="1260" w:id="3"/>') + run("：山田") + '</w:p>'
    '<w:p><w:pPr><w:jc w:val="distribute"/></w:pPr>' + run("均等割り付け") + '</w:p>'
    '<w:p><w:pPr><w:ind w:firstLineChars="100" w:leftChars="200"/></w:pPr>' + run("字数の字下げ") + '</w:p>'
    '<w:p><w:pPr><w:kinsoku w:val="0"/><w:overflowPunct w:val="0"/><w:topLinePunct w:val="1"/><w:autoSpaceDE w:val="0"/>'
    '<w:autoSpaceDN w:val="0"/><w:textAlignment w:val="center"/></w:pPr>' + run("段落の日本語の設定") + '</w:p>')
docx("06-japanese-properties.docx",
     body + sect('<w:textDirection w:val="tbRl"/><w:docGrid w:type="linesAndChars" w:linePitch="360" w:charSpace="-2000"/>'))

# Spreadsheets
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "和暦"
ws.column_dimensions["A"].width = 14
ws.column_dimensions["B"].width = 40
dates = [datetime.date(1868, 10, 23), datetime.date(1912, 7, 30), datetime.date(1926, 12, 25), datetime.date(1989, 1, 7),
         datetime.date(1989, 1, 8), datetime.date(2019, 4, 30), datetime.date(2019, 5, 1), datetime.date(2026, 10, 6)]
r = 1
for code in ['[$-411]ggge"年"m"月"d"日"', '[$-ja-JP-x-gannen]ggge"年"m"月"d"日";@', '[$-411]gge"."m"."d', '[$-411]ge"."m"."d',
             '[$-411]ggg ee"年"', '[$-411]yyyy"年"m"月"d"日"(aaa)', '[$-411]aaaa']:
    ws.cell(r, 1, code)
    r += 1
    for d in dates:
        ws.cell(r, 1, d.isoformat())
        ws.cell(r, 2, d).number_format = code
        r += 1
wb.save(f"{OUT}/07-wareki.xlsx")

wb = openpyxl.Workbook()
ws = wb.active
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 30
for i, code in enumerate(['[DBNum1][$-411]General', '[DBNum2][$-411]General', '[DBNum3][$-411]General', '[DBNum1][$-411]#,##0']):
    ws.cell(i + 1, 1, code)
    ws.cell(i + 1, 2, 1234).number_format = code
wb.save(f"{OUT}/08-kansuji.xlsx")

wb = openpyxl.Workbook()
ws = wb.active
rows = [("ｱｲｳ", "=ASC(A1)", "=UNICODE(ASC(A1))"), ("アイウ", "=ASC(A2)", None), ("ABC123ｱｲｳ", "=JIS(A3)", None),
        ("あいう", "=LENB(A4)", "=LEFTB(A4,2)"), ("１２３", "=VALUE(A5)", None), (1234, "=NUMBERSTRING(A6,1)", None),
        (datetime.date(2019, 5, 1), "=DATESTRING(A7)", None), ("山田", "=PHONETIC(A8)", None), (1234, "=YEN(A9)", None)]
for i, (a, b, c) in enumerate(rows, 1):
    ws.cell(i, 1, a)
    ws.cell(i, 2, b)
    if c:
        ws.cell(i, 3, c)
for col in "ABC":
    ws.column_dimensions[col].width = 24
wb.save(f"{OUT}/09-functions.xlsx")

wb = openpyxl.Workbook()
ws = wb.active
ws["B2"] = "縦書きＡＢ「題」ー"
ws["B2"].alignment = Alignment(text_rotation=255, vertical="top")
ws["B2"].font = Font(name="IPAexMincho", size=11)
ws.row_dimensions[2].height = 200
ws.column_dimensions["B"].width = 5
wb.save(f"{OUT}/10-cell-stacked.xlsx")

# Furigana in shared strings (written by hand: openpyxl keeps no rPh)
parts = {
    "[Content_Types].xml": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>',
    "_rels/.rels": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>',
    "xl/workbook.xml": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="名簿" sheetId="1" r:id="rId1"/></sheets></workbook>',
    "xl/_rels/workbook.xml.rels": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>',
    "xl/sharedStrings.xml": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="1" uniqueCount="1"><si><t>山田太郎</t><rPh sb="0" eb="2"><t>ヤマダ</t></rPh><rPh sb="2" eb="4"><t>タロウ</t></rPh><phoneticPr fontId="0" type="Hiragana"/></si></sst>',
    "xl/worksheets/sheet1.xml": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="s" ph="1"><v>0</v></c><c r="B1" t="str"><f>PHONETIC(A1)</f><v>やまだたろう</v></c></row></sheetData><phoneticPr fontId="0" type="Hiragana"/></worksheet>',
}
with zipfile.ZipFile(f"{OUT}/11-furigana.xlsx", "w", zipfile.ZIP_DEFLATED) as z:
    for k, v in parts.items():
        z.writestr(k, v)
