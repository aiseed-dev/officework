#!/usr/bin/env python3
"""**xlsx の式の答えを、Excel の計算と突き合わせる道具(macOS)。**

officework が書いた xlsx を Excel に開かせて計算させ、式のセルの値を
AppleScript で読み、officework 自身が計算した値とセルごとに比べます。

    .venv/bin/python tools/ms_values.py 台帳.xlsx [台帳2.xlsx …]

違いがあれば1行ずつ出し、無ければ「一致」と出します。
数は 1e-9 の相対誤差まで同じと見ます。日付は通し番号に戻して比べます。
"""
import datetime
import os
import subprocess
import sys

from openpyxl import load_workbook


def excel_values(src, coords):
    """Excel に開かせて計算させ、`coords`(シート名, セル番地)の値を字で返す。

    保存はしません — `save workbook as` は「保存した」と言いながらファイルを
    書かないことがありました(2026-09-08)。値は AppleScript で1つずつ読み、
    論理値は TRUE / FALSE、日付は 1899-12-30 起点の通し番号、数は AppleScript の
    字(15 桁)にします。開き方の注意は tools/ms_pdf.py と同じです
    """
    name = os.path.basename(src)
    subprocess.run(["osascript", "-e", f'''tell application "Microsoft Excel"
    try
        close workbook "{name}" saving no
    end try
end tell'''], capture_output=True, text=True, timeout=120)
    subprocess.run(["open", "-a", "Microsoft Excel", os.path.abspath(src)], check=True)
    reads = "\n".join(
        f'set vv to value of range "{c}" of sheet "{ws}" of wb\n'
        "if vv is true then\n set tt to \"TRUE\"\nelse if vv is false then\n set tt to \"FALSE\"\nelse\n"
        " try\n  set tt to ((vv - epoch) / 86400) as text\n on error\n  set tt to vv as text\n end try\nend if\n"
        "set outs to outs & tt & (ASCII character 31)"
        for ws, c in coords
    )
    script = f'''
with timeout of 600 seconds
tell application "Microsoft Excel"
    set display alerts to false
    set wb to missing value
    repeat 180 times
        try
            if (name of every workbook) contains "{name}" then set wb to workbook "{name}"
        end try
        if wb is not missing value then exit repeat
        delay 1
    end repeat
    if wb is missing value then error "Excel が {name} を開きませんでした"
    calculate
    set epoch to current date
    set year of epoch to 1899
    set month of epoch to 12
    set day of epoch to 30
    set time of epoch to 0
    set outs to ""
    {reads}
    close wb saving no
    outs
end tell
end timeout
'''
    r = subprocess.run(["osascript", "-e", script], capture_output=True, text=True)
    if r.returncode != 0:
        raise RuntimeError(r.stderr.strip())
    vals = r.stdout.rstrip("\n").split("\x1f")[:-1]
    if len(vals) != len(coords):
        raise RuntimeError(f"答えの数が合いません: {len(vals)} / {len(coords)}")
    return dict(zip(coords, vals))


def norm(v):
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, datetime.datetime):
        d = v - datetime.datetime(1899, 12, 30)
        return d.days + d.seconds / 86400
    if isinstance(v, datetime.date) and not isinstance(v, datetime.datetime):
        return (v - datetime.date(1899, 12, 30)).days
    if isinstance(v, datetime.time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    if isinstance(v, datetime.timedelta):
        return v.days + v.seconds / 86400
    if v is None:
        return ""
    return v


def same(a, b):
    """うちの値(数・字・論理値)と、Excel から読んだ字を比べる"""
    a = norm(a)
    if isinstance(a, bool):
        return b == ("TRUE" if a else "FALSE")
    if isinstance(a, (int, float)):
        try:
            return abs(a - float(b)) <= 1e-9 * max(1.0, abs(a), abs(float(b)))
        except ValueError:
            return False
    return str(a) == str(b)


def compare(src):
    from officework import sheet

    formulas = load_workbook(src)
    coords = [
        (ws.title, c.coordinate)
        for ws in formulas.worksheets
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    theirs = excel_values(src, coords)
    ours = sheet.Book.open(src)
    diffs = []
    for ws, c in coords:
        mine = ours[ws][c].value
        excel = theirs[(ws, c)]
        if not same(mine, excel):
            diffs.append(f"{ws}!{c} {formulas[ws][c].value}: うち {mine!r} / Excel {excel!r}")
    return len(coords), diffs


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    bad = 0
    for src in sys.argv[1:]:
        n, diffs = compare(src)
        print(f"== {os.path.basename(src)}: 式 {n} 個, 違い {len(diffs)} 件")
        for d in diffs:
            print("  " + d)
        bad += len(diffs)
    sys.exit(1 if bad else 0)
